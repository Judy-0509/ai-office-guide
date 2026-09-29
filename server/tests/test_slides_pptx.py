from __future__ import annotations

import io
import json

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from aioffice.slides import pptx_export, spec
from aioffice.slides.server import DEFAULT_DECK_PATH


def _sample_deck() -> dict:
    deck = json.loads(DEFAULT_DECK_PATH.read_text(encoding="utf-8"))
    deck["slides"] = [spec.validate_slide(s) for s in deck["slides"]]
    return deck


def _contains_korean(text: str) -> bool:
    return any("가" <= ch <= "힣" for ch in text)


def test_export_bytes_produces_a_valid_pptx_with_native_charts():
    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    slides = list(prs.slides)
    assert len(slides) == 3

    chart_types = [shape.chart.chart_type for slide in slides for shape in slide.shapes
                    if shape.has_chart]
    assert XL_CHART_TYPE.COLUMN_CLUSTERED in chart_types  # bar panel
    assert XL_CHART_TYPE.BAR_STACKED_100 in chart_types   # stacked panel
    assert XL_CHART_TYPE.XY_SCATTER in chart_types        # dots panels (2 of them)
    assert chart_types.count(XL_CHART_TYPE.XY_SCATTER) == 2


def test_export_has_a_native_table_with_the_right_cell_text():
    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    table_slide = list(prs.slides)[1]
    table_shape = next(s for s in table_slide.shapes if s.has_table)
    table = table_shape.table
    assert len(table.rows) == 9  # header + 8 institutions
    assert [table.cell(0, c).text for c in range(len(table.columns))] == \
        ["기관", "날짜", "시각", "핵심 주장", "핵심 수치"]
    assert table.cell(1, 0).text == "A증권"
    assert table.cell(1, 4).text == "215백만 대"


def test_export_keeps_korean_text_intact():
    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    found = False
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and _contains_korean(shape.text_frame.text):
                found = True
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if _contains_korean(cell.text):
                            found = True
    assert found


def test_export_slide_dimensions_match_the_1920x1080_canvas():
    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    assert prs.slide_width == pptx_export._emu(1920)
    assert prs.slide_height == pptx_export._emu(1080)


def test_export_to_file_writes_a_readable_pptx(tmp_path):
    out = tmp_path / "deck.pptx"
    pptx_export.export_to_file(_sample_deck(), out)
    assert out.exists()
    prs = Presentation(str(out))
    assert len(list(prs.slides)) == 3


# --- A: no chart title, chart text is 12pt/Malgun Gothic ------------------------------------
def test_charts_have_no_title_and_use_12pt_malgun_gothic():
    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    charts = [shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart]
    assert charts
    for chart in charts:
        assert chart.has_title is False
        assert chart.font.size == pptx_export.Pt(pptx_export.CHART_FONT_PT)
        assert chart.font.name == "Malgun Gothic"


# --- B: bar chart -----------------------------------------------------------------------
def test_bar_chart_hides_value_axis_and_uses_outside_end_labels():
    from pptx.enum.chart import XL_LABEL_POSITION

    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    slide0 = list(prs.slides)[0]
    bar_chart = next(s.chart for s in slide0.shapes
                      if s.has_chart and s.chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED)
    assert bar_chart.value_axis.visible is False
    assert bar_chart.value_axis.has_major_gridlines is False
    plot = bar_chart.plots[0]
    assert plot.gap_width == 80
    assert plot.data_labels.position == XL_LABEL_POSITION.OUTSIDE_END
    # full category names, never truncated
    categories = list(plot.categories)
    assert "A증권" in categories and "H리서치" in categories


# --- C: stacked chart shows the raw value + part name, never a bogus "x00%" -------------------
def test_stacked_chart_labels_show_raw_value_with_part_name_not_percent():
    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    slide0 = list(prs.slides)[0]
    stacked_chart = next(s.chart for s in slide0.shapes
                          if s.has_chart and s.chart.chart_type == XL_CHART_TYPE.BAR_STACKED_100)
    assert stacked_chart.has_legend is False
    label_texts = []
    for series in stacked_chart.plots[0].series:
        for point in series.points:
            assert point.data_label.has_text_frame
            label_texts.append(point.data_label.text_frame.text)
    assert "부정 6" in label_texts
    assert "중립 3" in label_texts
    assert "긍정 2" in label_texts
    assert not any("%" in t for t in label_texts)  # never "600%"


# --- D: dots/scatter chart -----------------------------------------------------------------
def test_scatter_chart_hides_axes_orders_top_to_bottom_and_labels_are_value_only():
    from pptx.enum.chart import XL_LABEL_POSITION

    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    slide0 = list(prs.slides)[0]
    scatter_chart = next(s.chart for s in slide0.shapes
                          if s.has_chart and s.chart.chart_type == XL_CHART_TYPE.XY_SCATTER)
    assert scatter_chart.value_axis.visible is False
    assert scatter_chart.value_axis.has_major_gridlines is False
    assert scatter_chart.category_axis.has_major_gridlines is True
    n = len(list(scatter_chart.plots[0].series))
    # item order top->bottom: the first series (A증권) must have the HIGHEST y (n-1), so it
    # plots at the top of the panel.
    first_series = scatter_chart.plots[0].series[0]
    y_values = list(first_series.values)
    assert y_values == [n - 1]
    for series in scatter_chart.plots[0].series:
        point = series.points[0]
        assert point.data_label.has_text_frame
        assert point.data_label.position == XL_LABEL_POSITION.RIGHT
        # value only -- the item name now lives in the separate left-column text box, not the
        # chart's own data label (a data label carrying both collided when rows sat close
        # together in a short panel; see the fix-round notes).
        assert series.name not in point.data_label.text_frame.text
        assert point.data_label.text_frame.text.strip()


def test_scatter_chart_has_a_manual_plot_area_layout():
    """The inner plot area is pinned to fixed fractions (c:manualLayout) so row pixel
    positions are deterministic and the external left-label column can line up with the dots
    exactly -- without this, PowerPoint auto-sizes the plot area from label content and the
    two would drift apart."""
    from pptx.oxml.ns import qn

    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    slide0 = list(prs.slides)[0]
    scatter_chart = next(s.chart for s in slide0.shapes
                          if s.has_chart and s.chart.chart_type == XL_CHART_TYPE.XY_SCATTER)
    plot_area = scatter_chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    manual_layout = plot_area.find(qn("c:layout")).find(qn("c:manualLayout"))
    assert manual_layout is not None
    assert manual_layout.find(qn("c:layoutTarget")).get("val") == "inner"
    assert manual_layout.find(qn("c:xMode")).get("val") == "edge"
    assert manual_layout.find(qn("c:yMode")).get("val") == "edge"
    for tag in ("c:x", "c:y", "c:w", "c:h"):
        val = float(manual_layout.find(qn(tag)).get("val"))
        assert 0.0 <= val <= 1.0


def test_scatter_chart_left_label_column_has_one_textbox_per_item_right_aligned():
    from pptx.enum.text import PP_ALIGN

    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    slide0 = list(prs.slides)[0]
    # the dots panel's item labels are plain textboxes (not part of any chart)
    label_boxes = [
        s for s in slide0.shapes
        if s.has_text_frame and not s.has_chart and not s.has_table
        and s.text_frame.text in ("A증권", "B증권", "C투자증권", "E증권", "G리서치", "H리서치")
    ]
    assert len(label_boxes) == 6
    for box in label_boxes:
        assert box.text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT


def test_scatter_chart_uses_a_uniform_font_size_for_a_tight_row_pitch():
    """A short (2-3 panel layout) dots panel has ~13px rows -- well under the 18pt/36px
    threshold -- so every label in that chart (not just the colliding pair) must use the
    smaller 10pt size."""
    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    slide0 = list(prs.slides)[0]
    scatter_chart = next(s.chart for s in slide0.shapes
                          if s.has_chart and s.chart.chart_type == XL_CHART_TYPE.XY_SCATTER)
    sizes = set()
    for series in scatter_chart.plots[0].series:
        for para in series.points[0].data_label.text_frame.paragraphs:
            for run in para.runs:
                sizes.add(run.font.size)
    assert sizes == {pptx_export.Pt(10)}


# --- F: table column widths are proportional to content, floored at min_frac -----------------
def test_table_column_widths_are_proportional_not_equal():
    widths = pptx_export._column_widths_px(
        ["기관", "핵심 주장"], [["A증권", "매우 길고 상세한 핵심 주장 문장입니다"]], total_w_px=1000)
    assert len(widths) == 2
    assert widths[1] > widths[0]  # the long column gets more room than the short one
    assert sum(widths) == 1000
    assert all(w >= 1000 * 0.08 for w in widths)  # the 8% floor holds even for short columns


# --- M: highlight_rows (a list, not a single index) -------------------------------------------
def test_table_highlights_multiple_rows_via_highlight_rows():
    prs = Presentation(io.BytesIO(pptx_export.export_bytes(_sample_deck())))
    table_slide = list(prs.slides)[1]
    table_shape = next(s for s in table_slide.shapes if s.has_table)
    table = table_shape.table
    # deck_memory.json's table panel sets "highlight_rows": [6] -- still works as a list of one
    tint = pptx_export._rgb(spec.ACCENT_TINT)
    assert table.cell(7, 0).fill.fore_color.rgb == tint  # row index 6 -> table row 7 (1 header row)
    assert table.cell(1, 0).fill.fore_color.rgb != tint
