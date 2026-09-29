from __future__ import annotations

import io
import json
import threading
import urllib.error
import urllib.request

import pytest
from pptx import Presentation

from aioffice.config import Settings
from aioffice.slides import server as slides_server

API_KEY = "sk-super-secret-test-key-do-not-leak"


# --- J: SLIDES_MAX_TOKENS resolution (process env wins, then the --env file, then default) ---
def test_resolve_slides_max_tokens_reads_the_env_file(tmp_path, monkeypatch):
    monkeypatch.delenv("SLIDES_MAX_TOKENS", raising=False)
    env_file = tmp_path / ".env"
    env_file.write_text("SLIDES_MAX_TOKENS=9000\n", encoding="utf-8")
    assert slides_server.resolve_slides_max_tokens(env_file) == 9000


def test_resolve_slides_max_tokens_process_env_wins_over_env_file(tmp_path, monkeypatch):
    env_file = tmp_path / ".env"
    env_file.write_text("SLIDES_MAX_TOKENS=9000\n", encoding="utf-8")
    monkeypatch.setenv("SLIDES_MAX_TOKENS", "12345")
    assert slides_server.resolve_slides_max_tokens(env_file) == 12345


def test_resolve_slides_max_tokens_defaults_when_unset(tmp_path, monkeypatch):
    monkeypatch.delenv("SLIDES_MAX_TOKENS", raising=False)
    assert slides_server.resolve_slides_max_tokens(None) == slides_server.DEFAULT_SLIDES_MAX_TOKENS
    assert slides_server.resolve_slides_max_tokens(tmp_path / "missing.env") == \
        slides_server.DEFAULT_SLIDES_MAX_TOKENS


class FakeLLMClient:
    """Stands in for aioffice.llm.client.LLMClient: no network, deterministic events."""

    def __init__(self, settings: Settings, events=None):
        self.settings = settings
        self._events = events if events is not None else []
        self._events_queue: list[list] | None = None  # one event-list per call, for retry tests
        self.calls: list[dict] = []

    def stream_complete(self, messages, *, step, agent, task_id=None, model=None,
                         max_tokens=None, cancel=None):
        self.calls.append({"messages": messages, "step": step, "max_tokens": max_tokens})
        if self._events_queue is not None:
            events = self._events_queue.pop(0) if self._events_queue else []
        else:
            events = self._events
        yield from events

    def estimate_tokens(self, text: str) -> int:
        return max(1, len(text) // 2)


def _table_slide(title="T"):
    return {"layout": "table", "kicker": "01", "title": title,
            "table": {"type": "table", "title": "표", "columns": ["a"], "rows": [["1"]]}}


@pytest.fixture()
def studio(tmp_path):
    settings = Settings.load(None, environ={
        "DATA_DIR": str(tmp_path / "data"),
        "LLM_API_KEY": API_KEY,
        "LLM_MODEL_DEFAULT": "qwen-test",
    })
    settings.ensure_dirs()
    srv = slides_server.build_server(
        settings, slides_server.DEFAULT_TOPIC_PATH, tmp_path / "studio", None, "127.0.0.1", 0)
    srv.llm_client = FakeLLMClient(settings)  # type: ignore[attr-defined]
    port = srv.server_address[1]
    thread = threading.Thread(target=srv.serve_forever, daemon=True)
    thread.start()
    try:
        yield srv, port
    finally:
        srv.shutdown()
        thread.join(5)
        srv.server_close()


def _get(port, path):
    with urllib.request.urlopen(f"http://127.0.0.1:{port}{path}") as r:
        return r.status, json.loads(r.read().decode("utf-8"))


def _request_json(port, path, method, payload):
    data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(f"http://127.0.0.1:{port}{path}", method=method, data=data,
                                  headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req) as r:
            return r.status, json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        return exc.code, json.loads(exc.read().decode("utf-8"))


def _post_sse(port, path, payload):
    data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(f"http://127.0.0.1:{port}{path}", method="POST", data=data,
                                  headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req) as r:
        raw = r.read().decode("utf-8")
    events = []
    for block in raw.split("\n\n"):
        if not block.strip():
            continue
        event_type, data_line = "message", ""
        for line in block.split("\n"):
            if line.startswith("event:"):
                event_type = line[len("event:"):].strip()
            elif line.startswith("data:"):
                data_line += line[len("data:"):].strip()
        events.append((event_type, json.loads(data_line) if data_line else None))
    return events


# --- /api/state -----------------------------------------------------------------------------
def test_state_returns_sample_deck_warnings_topic_and_model(studio):
    srv, port = studio
    status, data = _get(port, "/api/state")
    assert status == 200
    assert len(data["deck"]["slides"]) == 3
    assert data["warnings"] == [[], [], []]
    assert data["model"] == "qwen-test"
    assert data["busy"] is False


def test_index_serves_html(studio):
    srv, port = studio
    with urllib.request.urlopen(f"http://127.0.0.1:{port}/") as r:
        body = r.read().decode("utf-8")
        content_type = r.headers.get("Content-Type")
    assert "text/html" in content_type
    assert "<html" in body.lower()


# --- PUT /api/deck ----------------------------------------------------------------------
def test_put_deck_saves_and_writes_a_numbered_snapshot(studio):
    srv, port = studio
    _, state = _get(port, "/api/state")
    deck = state["deck"]
    deck["slides"][0]["title"] = "수정된 제목"
    status, data = _request_json(port, "/api/deck", "PUT", deck)
    assert status == 200
    assert data["deck"]["slides"][0]["title"] == "수정된 제목"

    data_dir = srv.store.data_dir
    saved = json.loads((data_dir / "deck.json").read_text(encoding="utf-8"))
    assert saved["slides"][0]["title"] == "수정된 제목"
    versions = sorted((data_dir / "versions").glob("*.json"))
    assert len(versions) >= 1
    assert json.loads(versions[-1].read_text(encoding="utf-8"))["slides"][0]["title"] == "수정된 제목"


def test_put_deck_rejects_a_structurally_invalid_slide(studio):
    srv, port = studio
    status, data = _request_json(port, "/api/deck", "PUT",
                                  {"title": "t", "footer": "f", "slides": [{"layout": "nope"}]})
    assert status == 400
    assert "error" in data


# --- generation SSE ------------------------------------------------------------------------
def test_generate_streams_reasoning_then_op_then_done(studio):
    srv, port = studio
    slide = _table_slide("새 슬라이드")
    srv.llm_client._events = [
        ("reasoning", "생각 중..."),
        ("content", json.dumps({"op": "insert", "index": 0, "slide": slide}, ensure_ascii=False) + "\n"),
        ("usage", {"prompt_tokens": 20, "completion_tokens": 10}),
    ]

    events = _post_sse(port, "/api/generate", {"slides": 1})

    types = [t for t, _ in events]
    assert types[0] == "reasoning"
    assert "op" in types
    assert types[-1] == "done"
    op_payload = next(p for t, p in events if t == "op")
    assert op_payload["op"] == "insert"
    assert op_payload["slide"]["title"] == "새 슬라이드"
    done_payload = next(p for t, p in events if t == "done")
    assert done_payload["prompt_tokens"] == 20
    assert done_payload["completion_tokens"] == 10
    assert done_payload["internal_estimate_seconds"] == round(10 / 40, 1)

    status, state = _get(port, "/api/state")
    assert len(state["deck"]["slides"]) == 1
    assert state["deck"]["slides"][0]["title"] == "새 슬라이드"


def test_generate_reports_invalid_op_lines_as_warnings_but_keeps_going(studio):
    srv, port = studio
    good = _table_slide("OK")
    srv.llm_client._events = [
        ("content", "이건 산문입니다\n"),
        ("content", json.dumps({"op": "insert", "index": 0, "slide": good}, ensure_ascii=False) + "\n"),
    ]
    events = _post_sse(port, "/api/generate", {"slides": 1})
    assert any(t == "warning" for t, _ in events)
    assert any(t == "op" for t, _ in events)


def test_generate_retries_once_on_zero_ops_then_succeeds(studio):
    """L: the model sometimes finishes still "thinking" and emits no op lines at all. One
    retry with the same messages should recover it -- and the client sees a warning, not
    silence, while that happens."""
    srv, port = studio
    slide = _table_slide("복구됨")
    srv.llm_client._events_queue = [
        [("reasoning", "생각만 하다가 끝남")],  # attempt 1: zero ops, no error
        [("content", json.dumps({"op": "insert", "index": 0, "slide": slide}, ensure_ascii=False) + "\n")],
    ]
    events = _post_sse(port, "/api/generate", {"slides": 1})
    assert len(srv.llm_client.calls) == 2
    warning_texts = [p["text"] for t, p in events if t == "warning"]
    assert any("한 번 더 시도" in w for w in warning_texts)
    assert not any(t == "error" for t, _ in events)
    op_payload = next(p for t, p in events if t == "op")
    assert op_payload["slide"]["title"] == "복구됨"


def test_generate_gives_up_with_an_error_after_two_zero_ops_attempts(studio):
    srv, port = studio
    srv.llm_client._events_queue = [
        [("reasoning", "생각만 함")],
        [("reasoning", "또 생각만 함")],
    ]
    events = _post_sse(port, "/api/generate", {"slides": 1})
    assert len(srv.llm_client.calls) == 2
    assert not any(t == "op" for t, _ in events)
    error_texts = [p["text"] for t, p in events if t == "error"]
    assert any("생성하지 못했습니다" in e for e in error_texts)


def test_generate_does_not_retry_when_the_stream_already_errored(studio):
    srv, port = studio
    srv.llm_client._events_queue = [[("error", "네트워크 오류")]]
    events = _post_sse(port, "/api/generate", {"slides": 1})
    assert len(srv.llm_client.calls) == 1  # a real error is not "silence" -- don't mask it with a retry
    assert not any("한 번 더 시도" in p.get("text", "") for t, p in events if t == "warning")


def test_generate_returns_409_while_busy(studio):
    srv, port = studio
    with srv.store.lock:
        srv.store.busy = True
    try:
        status, data = _request_json(port, "/api/generate", "POST", {"slides": 3})
        assert status == 409
        assert "생성 중" in data["error"]
    finally:
        with srv.store.lock:
            srv.store.busy = False


def test_edit_returns_409_while_busy(studio):
    srv, port = studio
    with srv.store.lock:
        srv.store.busy = True
    try:
        status, data = _request_json(port, "/api/edit", "POST",
                                      {"instruction": "제목 줄여줘", "selected": 0})
        assert status == 409
    finally:
        with srv.store.lock:
            srv.store.busy = False


def test_edit_passes_selected_index_1_based_to_the_model(studio):
    srv, port = studio
    slide = _table_slide("OK")
    srv.llm_client._events = [
        ("content", json.dumps({"op": "insert", "index": 0, "slide": slide}, ensure_ascii=False) + "\n"),
    ]
    _post_sse(port, "/api/edit", {"instruction": "제목 줄여줘", "selected": 1})
    assert len(srv.llm_client.calls) == 1  # produced an op, so no zero-ops retry (see L)
    user_message = srv.llm_client.calls[0]["messages"][1]["content"]
    assert "선택된 슬라이드 번호: 2" in user_message
    assert "제목 줄여줘" in user_message


# --- export ----------------------------------------------------------------------------
def test_export_pptx_returns_a_valid_presentation(studio):
    srv, port = studio
    with urllib.request.urlopen(f"http://127.0.0.1:{port}/api/export.pptx") as r:
        body = r.read()
        content_type = r.headers.get("Content-Type")
        disposition = r.headers.get("Content-Disposition")
    assert content_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert "weekly-coverage.pptx" in disposition
    prs = Presentation(io.BytesIO(body))
    assert len(list(prs.slides)) == 3


# --- deck/new + cancel -------------------------------------------------------------------
def test_deck_new_resets_to_the_sample_deck(studio):
    srv, port = studio
    with srv.store.lock:
        srv.store.deck = {"title": "x", "footer": "f", "slides": []}
        srv.store.save()
    status, data = _request_json(port, "/api/deck/new", "POST", {})
    assert status == 200
    assert len(data["deck"]["slides"]) == 3


def test_cancel_sets_the_current_cancel_event(studio):
    srv, port = studio
    with srv.store.lock:
        srv.store.cancel_event = threading.Event()
    status, data = _request_json(port, "/api/cancel", "POST", {})
    assert status == 200
    assert srv.store.cancel_event.is_set()


# --- the API key must never leak -----------------------------------------------------------
def test_api_key_never_appears_in_any_response(studio):
    srv, port = studio
    responses = []

    status, state = _get(port, "/api/state")
    responses.append(json.dumps(state, ensure_ascii=False))

    slide = _table_slide("T")
    srv.llm_client._events = [
        ("content", json.dumps({"op": "insert", "index": 0, "slide": slide}, ensure_ascii=False) + "\n"),
        ("error", "네트워크 오류: 인증 실패"),
    ]
    events = _post_sse(port, "/api/generate", {"slides": 1})
    responses.append(json.dumps(events, ensure_ascii=False))

    status, err = _request_json(port, "/api/deck", "PUT", {"title": "t", "footer": "f",
                                                             "slides": [{"layout": "nope"}]})
    responses.append(json.dumps(err, ensure_ascii=False))

    for body in responses:
        assert API_KEY not in body
